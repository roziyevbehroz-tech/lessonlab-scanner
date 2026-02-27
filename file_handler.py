import openpyxl
from docx import Document
import os

class FileHandler:
    def parse_excel(self, file_path):
        """Excel fayldan (.xlsx) testlarni o'qiydi."""
        wb = openpyxl.load_workbook(file_path)
        sheet = wb.active
        title = os.path.splitext(os.path.basename(file_path))[0]
        questions = []

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            question_text = str(row[0]).strip()
            correct_answer = str(row[-1]).strip()
            possible_options = row[1:5]

            # To'g'ri javob indexini aniqlash
            correct_idx = -1
            if correct_answer.isdigit():
                correct_idx = int(correct_answer) - 1
            elif correct_answer.lower() in 'abcd' and len(correct_answer) == 1:
                correct_idx = ord(correct_answer.lower()) - ord('a')

            options = []
            for idx, opt in enumerate(possible_options):
                if opt:
                    is_correct = (idx == correct_idx) or (str(opt).strip() == correct_answer)
                    options.append({'text': str(opt).strip(), 'is_correct': is_correct})

            if question_text and options:
                questions.append({'text': question_text, 'options': options})

        return {'title': title, 'questions': questions}

    def parse_word(self, file_path):
        """Word fayldan (.docx) matnni oladi."""
        doc = Document(file_path)
        return "\n".join(p.text for p in doc.paragraphs)
